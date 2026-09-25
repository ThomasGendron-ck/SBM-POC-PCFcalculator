"""Tests du pipeline « Fichier de collecte » : règles DQR et non-régression SORHOY15."""

from pathlib import Path

import pytest

from pcf_extraction.collect import (
    FLAG_PAS_DE_FE,
    compute_dqr,
)

pd = pytest.importorskip("pandas")

REPO = Path(__file__).resolve().parents[1]
INPUT_DIR = REPO / "input"

requires_data = pytest.mark.skipif(
    not INPUT_DIR.is_dir(),
    reason="dossier input absent du dépôt",
)


class TestComputeDqr:
    def test_geo_vide_ou_zero(self):
        assert compute_dqr(None, "CK_EF_Packaging", True)["RM GEO DQR"] == 5.0
        assert compute_dqr("", "CK_EF_Packaging", True)["RM GEO DQR"] == 5.0
        assert compute_dqr("0", "CK_EF_Packaging", True)["RM GEO DQR"] == 5.0

    def test_geo_glo_row_rer_autre(self):
        assert compute_dqr("GLO", "src", True)["RM GEO DQR"] == 4.0
        assert compute_dqr("RoW", "src", True)["RM GEO DQR"] == 3.0
        assert compute_dqr("RER", "src", True)["RM GEO DQR"] == 2.0
        assert compute_dqr("FR", "src", True)["RM GEO DQR"] == 1.0

    def test_tech_source(self):
        assert compute_dqr("FR", None, True)["RM TECH DQR"] == 5.0
        assert compute_dqr("FR", "EF PROXY bois", True)["RM TECH DQR"] == 3.0
        assert compute_dqr("FR", "CK_EF_Packaging", True)["RM TECH DQR"] == 1.0
        assert compute_dqr("FR", "EI3.10 SimaPro", True)["RM TECH DQR"] == 2.0

    def test_temp_dqr(self):
        assert compute_dqr("FR", "src", True)["RM TEMP DQR"] == 1.0
        assert compute_dqr(None, None, False)["RM TEMP DQR"] == 5.0

    def test_dqr_calc_moyenne(self):
        dqr = compute_dqr("RER", "EI3.10 SimaPro", True)
        assert dqr["RM DQR value"] == pytest.approx((2.0 + 2.0 + 1.0) / 3)
        dqr = compute_dqr(None, None, False)
        assert dqr["RM DQR value"] == pytest.approx(5.0)

    def test_pas_de_fe(self):
        assert compute_dqr(None, None, False)["RM DQR value"] == 5.0


@requires_data
class TestSorhoy15:
    """Non-régression sur l'exemple SORHOY15 du Fichier de collecte."""

    @pytest.fixture(scope="class")
    def collecte(self):
        from pcf_extraction.collect import build_collecte

        df = build_collecte(INPUT_DIR)
        return df[df["Product SKU"] == "SORHOY15"].reset_index(drop=True)

    def test_11_composants(self, collecte):
        assert len(collecte) == 11
        assert collecte["Component SKU"].nunique() == 11

    def test_pcf_value(self, collecte):
        pcf = collecte["PCF Value"].iloc[0]
        assert pcf == pytest.approx(0.686545, abs=1e-5)

    def test_dqr_product_exact(self, collecte):
        assert collecte["DQR Product"].iloc[0] == pytest.approx(1.582432719775901, abs=1e-9)

    def test_parts(self, collecte):
        parts = collecte["Part du composant dans le produit"]
        assert parts.sum() == pytest.approx(1.0, abs=1e-9)
        by_comp = collecte.set_index("Component SKU")["Part du composant dans le produit"]
        assert by_comp["200004"] == pytest.approx(0.0011556, abs=1e-6)
        assert by_comp["600232"] == pytest.approx(0.8667154, abs=1e-6)
        assert by_comp["555350FR"] == pytest.approx(0.0612479, abs=1e-6)

    def test_ghg_par_composant(self, collecte):
        by_comp = collecte.set_index("Component SKU")["RM GHG"]
        assert by_comp["407326"] == pytest.approx(0.025350, abs=1e-6)
        assert by_comp["403113"] == pytest.approx(0.001200, abs=1e-6)
        assert by_comp["513898"] == pytest.approx(0.001829, abs=1e-6)
        assert by_comp["600232"] == pytest.approx(0.570000, abs=1e-6)
        assert by_comp["402076"] == pytest.approx(0.008424, abs=1e-6)
        assert by_comp["402075"] == pytest.approx(0.008360, abs=1e-6)

    def test_composant_sans_fe(self, collecte):
        row = collecte[collecte["Component SKU"] == "200004"].iloc[0]
        assert pd.isna(row["RM EF Value"])
        assert FLAG_PAS_DE_FE in (row["Flag"] or "")

    def test_pds_toujours_nul(self, collecte):
        assert (collecte["RM PDS value"] == 0.0).all()
        assert (collecte["RM EF PDS"] == 0.0).all()


class TestFreight:
    """Bloc Impact Appro Transport : rattachement des trajets fret (v0.5)."""

    @pytest.fixture(scope="class")
    def collecte(self):
        from pcf_extraction.collect import build_collecte

        return build_collecte(INPUT_DIR)

    def test_colonnes_fret_presentes(self, collecte):
        for col in ("Freight Supplier Code", "Freight Supplier Name", "Freight Route", "Freight Transportation Mode", "Freight GHG", "Freight GHG Unit"):
            assert col in collecte.columns

    def test_trajets_renseignes(self, collecte):
        n = int(collecte["Freight Route"].notna().sum())
        assert n > 0


    def test_ghg_transport_coherence_impl(self, collecte):
        sub = collecte[collecte["Freight GHG"].notna()]
        assert len(sub) > 0
        assert (sub["Freight GHG"] >= 0).all()

    def test_flag_trajet_manquant(self, collecte):
        from pcf_extraction.collect import FLAG_PAS_DE_TRAJET_FRET

        sans = collecte[collecte["Freight Route"].isna() & collecte["Component SKU"].notna()]
        if not sans.empty:
            assert sans["Flag"].fillna("").str.contains(FLAG_PAS_DE_TRAJET_FRET).all()


class TestTransformation:
    """Bloc Impact fabrication fournisseur : formule Transformation GHG (règles v0.7)."""

    def test_ghg_transformation_sans_saisie(self):
        from pcf_extraction.collect import _compute_transformation_ghg

        df = pd.DataFrame(
            {
                "Transformation Process EF Value": [1.2, None],
                "Quantity": [2.0, 3.0],
                "Net Weight": [0.5, 0.1],
            }
        )
        out = _compute_transformation_ghg(df)
        assert out["Transformation GHG"].iloc[0] == pytest.approx(1.2)
        assert pd.isna(out["Transformation GHG"].iloc[1])

    def test_transformation_priorite_saisie(self, tmp_path):
        from pcf_extraction.collect import _apply_transformation

        collecte = pd.DataFrame(
            {
                "Product SKU": ["P1", "P1"],
                "Component SKU": ["C1", "C2"],
                "Transformation Process EF Value": [None, None],
                "Quantity": [1.0, 2.0],
                "Net Weight": [1.0, 1.0],
                "Transformation GHG": [None, None],
                "Transformation Process Name": [None, "Extrusion"],
                "Transformation Energy Name": [None, None],
                "Transformation Process EF Unit": [None, None],
                "Transformation EF Source": [None, None],
            }
        )
        saisie = pd.DataFrame(
            {
                "Product SKU": ["P1"],
                "Component SKU": ["C1"],
                "Transformation Process EF Value": [0.5],
                "Transformation Process Name": ["Injection"],
                "Transformation Energy Name": ["Électricité"],
                "Transformation Process EF Unit": ["kgCO2e/kg"],
                "Transformation EF Source": ["Fiche technique"],
            }
        )
        out = _apply_transformation(collecte, saisie)
        row = out[out["Component SKU"] == "C1"].iloc[0]
        assert row["Transformation GHG"] == pytest.approx(0.5)
        assert row["Transformation Process Name"] == "Injection"
        row2 = out[out["Component SKU"] == "C2"].iloc[0]
        assert row2["Transformation Process Name"] == "Extrusion"


class TestStyleCollecte:
    """Couleurs de blocs et groupements de colonnes (Plans > Grouper) dans le livrable."""

    def test_ordre_colonnes_v07(self):
        from pcf_extraction import config
        cols = [c for _, block in config.BLOCK_HEADERS for c in block]
        assert cols[2:6] == ["PCF Value", "PCF Unit", "DQR Product", "PDS Product"]
        assert cols.index("RM PDS Activity Data") > cols.index("Quantity")
        assert cols.index("RM PDS Activity Data") < cols.index("Net Weight")
        assert "Transformation Process EF Value" in cols
        assert "Transformation Energy EF Value" in cols
        assert "Transformation EF" not in cols

    def test_layout_v07(self):
        from pcf_extraction import config
        assert config.COLLECTE_HEADER_ROW == 6
        assert config.COLLECTE_FREEZE_PANES == "C7"

    def test_marqueurs_grouping_v071(self):
        from pcf_extraction import config
        cols = [c for _, block in config.BLOCK_HEADERS for c in block]
        for marker in config.GROUPING_MARKER_COLUMNS:
            assert marker in cols, marker
        assert config.EXCEL_HEADER_MAP["Component Material (2)"] == "Component Material"
        assert config.EXCEL_HEADER_MAP["Transformation Details (2)"] == "Transformation Details"

    def test_groupe_transformation_v071(self):
        from pcf_extraction import config
        cols = [c for _, block in config.BLOCK_HEADERS for c in block]
        first, last = ("Transformation Process Name", "Transformation Energy EF Source")
        assert cols.index(first) < cols.index(last)

    def test_layout_v072(self):
        from pcf_extraction import config
        assert config.GROUPING_MARKER_WIDTH == 9.43
        assert config.COLLECTE_HEADER_ROW_HEIGHT == 108.75
        assert config.COLLECTE_TOP_GROUP_ROWS == (1, 4)
        assert config.FLAG_FILL == "FFFFC000"

    def test_font_colors_v072(self):
        from pcf_extraction import config
        assert config.BLOCK_FONT_COLORS["Produit"] == "FFFFFFFF"
        assert config.BLOCK_FONT_COLORS["Composants"] == "FFFFFFFF"
        assert config.BLOCK_FONT_COLORS["Impact matière fournisseur"] == "FF000000"
        assert config.BLOCK_FONT_COLORS["Flag"] == "FF000000"

    def test_synthese_v072(self):
        from pcf_extraction.report import build_synthese, SYNTHESE_COLUMNS
        import pandas as pd
        collecte = pd.DataFrame(
            {
                "Product SKU": ["P1", "P1"],
                "Product Designation": ["Prod 1", "Prod 1"],
                "PCF Value": [1.5, 1.5],
                "DQR Product": [2.0, 2.0],
                "PDS Product": [0.0, 0.0],
                "Carbon category": ["Raw Material", "PACKAGING"],
                "RM GHG": [1.0, 0.5],
                "Transformation GHG": [None, None],
                "Freight GHG": [0.01, None],
                "RM PDS Activity Data": [1.0, 1.0],
                "RM EF PDS": [0.0, 0.0],
                "RM DQR value": [3.0, 2.0],
            }
        )
        syn = build_synthese(collecte)
        assert list(syn.columns) == SYNTHESE_COLUMNS
        row = syn.iloc[0]
        assert row["PCF Value"] == 1.5
        assert row["Raw Material - GHG Value"] == 1.0
        assert row["Packaging - GHG Value"] == 0.5
        assert row["Freight - GHG Value"] == 0.01
        assert row["Raw Material - % PCF total"] == pytest.approx(1.0 / 1.5)
        assert pd.isna(row["Use - GHG Value"])
        assert pd.isna(row["End of Life - GHG Value"])
        assert "Use - GHG Value" in syn.columns and "End of Life - GHG Value" in syn.columns

    def test_couleurs_v071(self):
        from pcf_extraction import config
        assert config.BLOCK_COLORS["Produit"] == "FF156082"
        assert config.BLOCK_COLORS["Produit PCF"] == "FF0A3041"
        assert config.BLOCK_COLORS["Composants"] == "FF13501B"
        assert config.BLOCK_COLORS["Impact Appro Transport"] == "FF6FC5E6"

    def test_formats_nombre(self):
        from pcf_extraction import config
        cols = [c for _, block in config.BLOCK_HEADERS for c in block]
        for col in config.NUMBER_FORMATS:
            assert col in cols, col
        assert config.NUMBER_FORMATS["RM EF Value"] == "0.00000"
        assert config.NUMBER_FORMATS["RM GHG"] == "0.000"
        assert config.NUMBER_FORMATS["Recycled %"] == "0"
        assert config.NUMBER_FORMATS["RM PDS Activity Data"] == "0"

    def test_structure_blocs(self):
        from pcf_extraction import config

        cols = [c for _, block in config.BLOCK_HEADERS for c in block]
        from pcf_extraction.collect import COLLECTE_COLUMNS

        assert cols == COLLECTE_COLUMNS

    def test_groupes_dans_limites(self):
        from pcf_extraction import config

        cols = [c for _, block in config.BLOCK_HEADERS for c in block]
        for first, last in config.COLUMN_GROUPS:
            assert first in cols and last in cols
            assert cols.index(first) <= cols.index(last)


class TestV073:
    """v0.73 : Synthèse en 1er onglet, formats décimaux/pourcentages, réinjection des FE ecoinvent."""

    def test_synthese_header_layout(self):
        from pcf_extraction import report
        assert report.SYNTHESE_HEADER_HEIGHT == 60.0

    def test_synthese_formats_v073(self, tmp_path):
        import pandas as pd
        from pcf_extraction.collect import COLLECTE_COLUMNS
        from pcf_extraction.report import write_collecte_report

        collecte = pd.DataFrame(
            {
                "Product SKU": ["P1", "P1"],
                "Product Designation": ["Prod 1", "Prod 1"],
                "PCF Value": [1.5, 1.5],
                "DQR Product": [2.0, 2.0],
                "PDS Product": [0.0, 0.0],
                "Carbon category": ["Raw Material", "PACKAGING"],
                "RM GHG": [1.0, 0.5],
                "Transformation GHG": [None, None],
                "Freight GHG": [0.01, None],
                "RM PDS Activity Data": [1.0, 1.0],
                "RM EF PDS": [0.0, 0.0],
                "RM DQR value": [3.0, 2.0],
                "RM EF Name": ["FE A", "FE B"],
                "Raw Material": ["Mat A", "Mat B"],
                "Component SKU": ["C1", "C2"],
                "Component Designation": ["Comp 1", "Comp 2"],
                "Flag": ["", ""],
            }
        )
        for col in COLLECTE_COLUMNS:
            if col not in collecte.columns:
                collecte[col] = None
        out = tmp_path / "collecte_v073.xlsx"
        write_collecte_report(collecte, out)
        from openpyxl import load_workbook

        wb = load_workbook(out)
        assert wb.sheetnames[0] == "Synthese"
        assert "Fichier de collecte" in wb.sheetnames
        ws = wb["Synthese"]
        assert ws.row_dimensions[1].height == 60.0
        assert ws.cell(row=1, column=1).alignment.wrap_text
        headers = [c.value for c in ws[1]]
        for idx, header in enumerate(headers, start=1):
            fmt = ws.cell(row=2, column=idx).number_format
            if header and header.endswith("% PCF total"):
                assert fmt == "0.00%", header
            elif header in ("Product SKU", "Product Designation"):
                assert fmt == "General", header
            else:
                assert fmt == "#,##0.000", header

    def test_build_fe_overrides(self):
        import pandas as pd
        from pcf_extraction.collect import build_fe_overrides

        matching = pd.DataFrame(
            {
                "Component SKU": ["100", "200"],
                "Dataset ecoinvent": ["market for fertiliser", "market for plastic"],
                "FE proposé (kg CO2e/kg)": [1.5, 2.5],
                "Géographie": ["FR", "RoW"],
                "Statut": ["MATCHÉ", "NON MATCHÉ"],
            }
        )
        overrides = build_fe_overrides(matching)
        assert set(overrides) == {"100"}
        o = overrides["100"]
        assert o["name"] == "market for fertiliser"
        assert o["value"] == 1.5
        assert o["geo"] == "FR"
        assert o["unit"] == "kgCO2e/kg"
        assert "EcoInvent 3.12 cut-off" in o["source"] and "validé SBM" in o["source"]

    @requires_data
    def test_reinjection_fe_ecoinvent(self):
        from pcf_extraction.collect import build_collecte, build_fe_overrides
        from pcf_extraction.ecoinvent import load_lcia_gwp, match_missing_fe

        lcia_files = list(INPUT_DIR.glob("Cut-off Cumulative LCIA*.xlsx"))
        if not lcia_files:
            pytest.skip("fichier LCIA absent du dossier input")
        collecte_sans = build_collecte(INPUT_DIR)
        matching = match_missing_fe(collecte_sans, load_lcia_gwp(lcia_files[0]))
        overrides = build_fe_overrides(matching)
        collecte_avec = build_collecte(INPUT_DIR, fe_overrides=overrides)
        sans_fe = collecte_sans["Flag"].astype(str).str.contains(FLAG_PAS_DE_FE, na=False)
        avec_fe = collecte_avec["Flag"].astype(str).str.contains(FLAG_PAS_DE_FE, na=False)
        assert avec_fe.sum() < sans_fe.sum()
        eco = collecte_avec["RM EF Source"].astype(str).str.contains("EcoInvent", na=False)
        assert eco.sum() > 0


class TestV074:
    """v0.74 : périmètre BOM élargi (BOMALT sauf 2 et 9, USESTA_0 = 2) et taux de perte/rebuts."""

    def test_filtre_bom(self):
        from pcf_extraction.collect import BOM_EXCLUDED_ALTERNATIVES, BOM_ACTIVE_STATUS
        assert BOM_EXCLUDED_ALTERNATIVES == (2, 9)
        assert BOM_ACTIVE_STATUS == 2

    def test_scrap_rate_dans_colonnes(self):
        from pcf_extraction import config
        cols = [c for _, block in config.BLOCK_HEADERS for c in block]
        assert "Scrap Rate" in cols
        assert config.NUMBER_FORMATS["Scrap Rate"] == "0.00000"
        assert "Scrap Rate" in config.SAISIE_COLUMNS
        assert any(col == "Scrap Rate" for col, _ in config.SAISIE_GUIDE_ROWS)

    def test_scrap_rate_applique_au_ghg_transformation(self):
        import numpy as np
        import pandas as pd
        from pcf_extraction.collect import _compute_transformation_ghg

        collecte = pd.DataFrame(
            {
                "Transformation Process EF Value": [2.0, 2.0],
                "Quantity": [2.0, 2.0],
                "Net Weight": [3.0, 3.0],
                "Scrap Rate": [0.05, np.nan],
            }
        )
        out = _compute_transformation_ghg(collecte)
        assert out["Transformation GHG"].iloc[0] == pytest.approx(2.0 * 2.0 * 3.0 * 1.05)
        assert out["Transformation GHG"].iloc[1] == pytest.approx(2.0 * 2.0 * 3.0)

    @requires_data
    def test_perimetre_bom_elargi(self):
        from pcf_extraction.collect import build_collecte

        collecte = build_collecte(INPUT_DIR)
        pcf = collecte.groupby("Product SKU")["PCF Value"].first()
        assert int(pcf.notna().sum()) >= 200
        sorhoy = collecte[collecte["Product SKU"] == "SORHOY15"]
        assert len(sorhoy) == 11
        assert sorhoy["PCF Value"].iloc[0] == pytest.approx(0.686545, abs=1e-5)
