"""Tests of the pcf_results.xlsx writer (spec v0.97)."""
import pandas as pd
import pytest
from openpyxl import load_workbook

from sbm_pcf_calculation.results_spec import MANDATORY_FIELDS, MISSING_EF_COLUMNS, PCF_COLUMNS
from sbm_pcf_calculation.results_writer import (
    build_missing_ef_sheet,
    build_pcf_sheet,
    write_pcf_results,
)


@pytest.fixture
def component_lines() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Product SKU": "SORHOY15",
                "Product Designation": "Spray niche",
                "PCF Value": 0.693,
                "PCF Unit": "kgCO2e/component",
                "DQR Product": 1.579,
                "PDS Product": 0.0,
                "Product Category Code": "PDTET",
                "Product Category description": "Etiquettes",
                "Pack Unit Box": 12.0,
                "Product Net Weight": 0.32,
                "Product Gross Weight": 0.35,
                "Product Supplier Code": "SUP001",
                "Product Supplier Name": "Supplier A",
                "Component SKU": "COMP1",
                "Component Designation": "Bottle",
                "Category Code": "PDTMA",
                "Category description": "RAW MATERIALS ACTIV INGREDIENT",
                "Carbon category": "Raw Material",
                "Pack unit box": 6.0,
                "Quantity": 2.0,
                "RM PDS Activity Data": 1.0,
                "Net Weight": 0.12,
                "Net Weight Unit": "kg",
                "Gross Weight": 0.14,
                "Gross Weight Unit": "kg",
                "Stock unit": "UNIT",
                "Supplier code": "SUP002",
                "Supplier Name": "Supplier B",
                "Raw Material": "PET",
                "Recycled %": 30,
                "RM EF Name": "PET FE",
                "RM EF Value": 2.5,
                "RM EF Unit": "kgCO2e/kg",
                "RM EF Source": "EcoInvent",
                "RM EF PDS": 0.0,
                "Prod_EF_Geography": "RER",
                "RM GEO DQR": 2.0,
                "RM TECH DQR": 2.0,
                "RM TEMP DQR": 1.0,
                "RM DQR value": 1.667,
                "RM GHG": 0.6,
                "Transformation GHG": 0.09,
                "Freight Supplier Code": "SUP002",
                "Freight Supplier Name": "Supplier B",
                "Freight Route": "ROUTE1",
                "Freight Transportation Mode": "TRUCK",
                "Freight GHG": 0.003,
                "Flag": "",
            },
            {
                "Product SKU": "SORHOY15",
                "Product Designation": "Spray niche",
                "PCF Value": 0.693,
                "PCF Unit": "kgCO2e/component",
                "DQR Product": 1.579,
                "PDS Product": 0.0,
                "Product Category Code": "PDTET",
                "Product Category description": "Etiquettes",
                "Pack Unit Box": 12.0,
                "Product Net Weight": 0.32,
                "Product Gross Weight": 0.35,
                "Product Supplier Code": "SUP001",
                "Product Supplier Name": "Supplier A",
                "Component SKU": "COMP2",
                "Component Designation": "Cap",
                "Category Code": "PDTEM",
                "Category description": "Packaging others materials",
                "Carbon category": "PACKAGING",
                "Pack unit box": 100.0,
                "Quantity": 1.0,
                "RM PDS Activity Data": 1.0,
                "Net Weight": 0.01,
                "Net Weight Unit": "kg",
                "Gross Weight": 0.01,
                "Gross Weight Unit": "kg",
                "Stock unit": "CARTON",
                "Supplier code": "SUP003",
                "Supplier Name": "Supplier C",
                "Raw Material": "PP",
                "Recycled %": 0,
                "RM EF Name": None,
                "RM EF Value": None,
                "RM EF Source": None,
                "RM GHG": None,
                "Freight GHG": None,
                "Flag": "PAS DE FE ASSOCIE A LA MATIERE",
            },
        ]
    )


@pytest.fixture
def product_results() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Product SKU": "SORHOY15",
                "Product Designation": "Spray niche",
                "PCF Value": 0.693,
                "DQR": 1.579,
                "PDS": 0.0,
                "Flag level": "HIGH",
            }
        ]
    )


class TestBuildPcfSheet:
    def test_column_order_matches_spec(self, component_lines):
        sheet = build_pcf_sheet(component_lines)
        assert list(sheet.columns) == [name for name, _ in PCF_COLUMNS]

    def test_renamed_columns_are_present(self, component_lines):
        sheet = build_pcf_sheet(component_lines)
        assert "PCF GHG Value" in sheet.columns
        assert "Component Category Code" in sheet.columns
        assert sheet.loc[0, "PCF GHG Value"] == 0.693
        assert sheet.loc[0, "Component Category Code"] == "PDTMA"

    def test_default_values_are_filled(self, component_lines):
        sheet = build_pcf_sheet(component_lines)
        assert sheet.loc[0, "RM EF PDS"] == 0
        assert sheet.loc[0, "Supplier PCF external review"] == "No"
        assert sheet.loc[0, "RM GHG Unit"] == "kgCO2e/component"

    def test_flag_level_mapped_to_validation_flag(self, component_lines, product_results):
        sheet = build_pcf_sheet(component_lines, product_results)
        assert (sheet["Data validation flag"] == "HIGH").all()


class TestBuildMissingEfSheet:
    def test_only_components_without_ef(self, component_lines):
        sheet = build_missing_ef_sheet(component_lines)
        assert list(sheet["Component SKU"]) == ["COMP2"]

    def test_column_order_matches_spec(self, component_lines):
        sheet = build_missing_ef_sheet(component_lines)
        assert list(sheet.columns) == [name for name, _ in MISSING_EF_COLUMNS]

    def test_matching_columns_mapped(self, component_lines):
        matching = pd.DataFrame(
            [
                {
                    "Component SKU": "COMP2",
                    "Dataset ecoinvent": "market for polypropylene",
                    "Règle de matching": "matière PP",
                    "FE proposé (kg CO2e/kg)": 1.8,
                    "Géographie": "RER",
                    "Statut": "MATCHÉ",
                }
            ]
        )
        sheet = build_missing_ef_sheet(component_lines, matching)
        row = sheet.iloc[0]
        assert row["RM AutoMatch EF Name"] == "market for polypropylene"
        assert row["RM AutoMatch EF Rationale"] == "matière PP"
        assert row["RM AutoMatch EF Value"] == 1.8
        assert row["RM AutoMatch EF Geography"] == "RER"


class TestWritePcfResults:
    def test_file_structure(self, tmp_path, component_lines, product_results):
        out = tmp_path / "pcf_results.xlsx"
        write_pcf_results(out, component_lines, product_results)
        workbook = load_workbook(out)
        assert workbook.sheetnames == ["PCF_Calculation", "MissingEF_Matching"]
        pcf_ws = workbook["PCF_Calculation"]
        assert pcf_ws.freeze_panes == "C3"
        missing_ws = workbook["MissingEF_Matching"]
        assert missing_ws.freeze_panes == "C3"

    def test_mandatory_row_and_header(self, tmp_path, component_lines, product_results):
        out = tmp_path / "pcf_results.xlsx"
        write_pcf_results(out, component_lines, product_results)
        workbook = load_workbook(out)
        missing_ws = workbook["MissingEF_Matching"]
        header = [cell.value for cell in missing_ws[2]]
        assert header[:2] == ["Component SKU", "Component Designation"]
        mandatory = [cell.value for cell in missing_ws[1]]
        mandatory_names = [
            header[i] for i, value in enumerate(mandatory) if value == "Mandatory"
        ]
        assert mandatory_names == MANDATORY_FIELDS

    def test_number_formats(self, tmp_path, component_lines, product_results):
        out = tmp_path / "pcf_results.xlsx"
        write_pcf_results(out, component_lines, product_results)
        workbook = load_workbook(out)
        pcf_ws = workbook["PCF_Calculation"]
        pcf_value_idx = [name for name, _ in PCF_COLUMNS].index("PCF GHG Value") + 1
        assert pcf_ws.cell(row=3, column=pcf_value_idx).number_format == "0.000"
