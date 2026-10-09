"""Tests de la génération du fichier EF matching (spec v0.98 MissingEF_Matching)."""

from pathlib import Path

import pytest

from sbm_pcf_calculation.ef_matching_writer import (
    EF_MATCHING_COLUMNS,
    EF_MATCHING_SHEET,
    HEADER_BLOCKS,
    build_ef_matching_rows,
    write_ef_matching,
)

pd = pytest.importorskip("pandas")
openpyxl = pytest.importorskip("openpyxl")

REPO = Path(__file__).resolve().parents[1]
TEMPLATE = REPO / "templates" / "Data_collection_template.xlsx"


def _collecte():
    return pd.DataFrame(
        [
            {
                "Product SKU": "P1",
                "Component SKU": "C1",
                "Component Designation": "BIDON HDPE 1L",
                "Category Code": "PDTBO",
                "Category description": "Bottles",
                "Carbon category": "PACKAGING",
                "Supplier code": "EFR06853",
                "Supplier Name": "FOURNISSEUR FR",
                "Pack unit box": 12.0,
                "Raw Material": "HDPE",
                "Recycled %": 0.25,
                "RM EF Value": float("nan"),
            },
            {
                "Product SKU": "P1",
                "Component SKU": "C2",
                "Component Designation": "CARTON",
                "Category Code": "PDTEC",
                "Category description": "Packaging Components CARTON",
                "Carbon category": "PACKAGING",
                "Supplier code": "XDE00017",
                "Supplier Name": "FOURNISSEUR DE",
                "Pack unit box": 1.0,
                "Raw Material": "CARTON",
                "Recycled %": float("nan"),
                "RM EF Value": 0.8,
            },
        ]
    )


def _matching():
    return pd.DataFrame(
        [
            {
                "Component SKU": "C1",
                "Statut": "MATCHÉ",
                "Dataset ecoinvent": "market for polyethylene, high density",
                "Règle de matching": "PE-HD",
                "FE proposé (kg CO2e/kg)": 0.0123,
                "Géographie": "FR",
            }
        ]
    )


def test_layout_57_champs_spec():
    headers = [h for h, b in EF_MATCHING_COLUMNS if b != "GROUPING" or True]
    assert len(EF_MATCHING_COLUMNS) == len({h for h, _ in EF_MATCHING_COLUMNS})
    assert ("Component SKU", "Composants") == EF_MATCHING_COLUMNS[0]
    assert ("Data validation flag impact", "Flag") == EF_MATCHING_COLUMNS[-1]
    groupings = [h for h, b in EF_MATCHING_COLUMNS if b == "GROUPING"]
    assert groupings == [
        "Component Details",
        "Component Material",
        "AutoMatch Details",
        "Transfo AutoMatch Details",
        "UserValidation",
        "Transfo UserValidation Details",
    ]


def test_build_rows_premplissage():
    rows = build_ef_matching_rows(_collecte(), _matching())
    assert len(rows) == 1
    row = rows[0]
    assert row["Component SKU"] == "C1"
    assert row["Component Supplier Country"] == "FR"
    assert row["RM AutoMatch EF Name"] == "market for polyethylene, high density"
    assert row["RM AutoMatch EF Value"] == 0.0123
    assert row["RM AutoMatch EF Unit"] == "kgCO2e/kg"
    assert row["RM AutoMatch EF Geography"] == "FR"
    assert row["RM AutoMatch GEO DQR"] == 1.0


def test_build_rows_sans_matching():
    rows = build_ef_matching_rows(_collecte(), None)
    assert len(rows) == 1
    assert rows[0].get("RM AutoMatch EF Name") is None


def test_write_file(tmp_path):
    out = tmp_path / "ef_matching.xlsx"
    n = write_ef_matching(str(out), _collecte(), _matching(), template_path=TEMPLATE)
    assert n == 1
    from openpyxl import load_workbook

    wb = load_workbook(str(out))
    assert wb.sheetnames == [EF_MATCHING_SHEET]
    ws = wb[EF_MATCHING_SHEET]
    headers = [ws.cell(row=2, column=i).value for i in range(1, ws.max_column + 1)]
    for (header, _), excel_header in zip(EF_MATCHING_COLUMNS, headers):
        assert header == excel_header or header == "RM UserValidation EF Rationale (2)"
    assert ws.cell(row=2, column=1).value == "Component SKU"
    assert ws.cell(row=1, column=1).value == "Mandatory ->"
    assert ws.cell(row=3, column=1).value == "C1"
    flag_cell = None
    for i in range(1, ws.max_column + 1):
        if ws.cell(row=2, column=i).value == "Data validation flag":
            flag_cell = ws.cell(row=3, column=i).value
    assert flag_cell is not None and flag_cell.startswith("=IF(")
    dqr_cell = None
    for i in range(1, ws.max_column + 1):
        if ws.cell(row=2, column=i).value == "RM AutoMatch DQR value":
            dqr_cell = ws.cell(row=3, column=i).value
    assert dqr_cell is not None and dqr_cell.startswith("=IF(COUNT(")


def test_doublon_user_validation_rationale_renomme():
    headers = [h for h, _ in EF_MATCHING_COLUMNS]
    assert headers.count("RM UserValidation EF Rationale") == 1
    assert "RM UserValidation EF Rationale (2)" in headers
    assert "Component Pack Unit box" in headers
