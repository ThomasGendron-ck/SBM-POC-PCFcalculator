"""Tests de la génération du fichier de collecte LM v0.95 (datashare)."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from pcf_extraction.datashare import (
    BLOCK_COLORS,
    DQR_GUIDE_SHEET,
    MANDATORY,
    PRODUCT_COLUMNS,
    SHEET_ORDER,
    TAB_COLORS,
    apply_v093,
    block_of,
)

INPUT = Path(__file__).resolve().parents[1] / "input" / "Data_collection_v091.xlsx"
TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "Data_collection_template.xlsx"


def _make_workbook(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb["Sheet"]
    ws.title = "Product"
    headers = [
        "Product SKU",
        "Supplier Name",
        "Net Weight",
        "Net Weight Unit",
        "Transformation Energy Name",
        "Transformation Process Scrap Rate",
        "Data validation flag",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h)
    ws.cell(row=2, column=1, value="SORHOY15")
    ws.cell(row=2, column=3, value=1.25)
    ws.cell(row=2, column=5, value="electricity")
    ws.cell(row=2, column=6, value=0.03)
    ug = wb.create_sheet("UserGuide")
    ug.cell(row=1, column=2, value="SBM - PCF - LM references - Data collection")
    path = tmp_path / "in.xlsx"
    wb.save(path)
    return path


def test_block_of():
    assert block_of("Product SKU", "Product") == "Produit"
    assert block_of("Transformation GHG", "Product") == "Transformation"
    assert block_of("Supplier PCF value", "Component") == "Supplier PCF"
    assert block_of("Supplier DQR", "Component") == "Supplier PCF"
    assert block_of("Supplier PCF framework", "Product") == "Supplier PCF"
    assert block_of("Data validation flag", "Component") == "Flag"
    assert block_of("Component SKU", "Component") == "Composants"
    assert block_of("Net Weight", "Component") == "Description du composant"
    assert block_of("Raw Material (MB Product)", "Component") == "Description du composant"


def test_columns_match_spec_v093():
    product_headers = [h for h, _ in PRODUCT_COLUMNS]
    assert product_headers.index("Supplier PCF framework") < product_headers.index("Supplier PCF value")
    assert product_headers.index("Supplier PCF declared unit") < product_headers.index("Supplier PCF value")
    assert product_headers.index("Transformation Energy Type") < product_headers.index("Transformation Process Name")
    assert product_headers.index("Transformation Process Comment") > product_headers.index("Transformation Process Scrap Rate")
    assert product_headers[-2:] == ["Data validation flag", "Data validation flag impact"]
    assert "Supplier PCF" in product_headers and "Transformation Details" in product_headers


def test_generate_from_synthetic(tmp_path):
    in_path = _make_workbook(tmp_path)
    out_path = tmp_path / "out.xlsx"
    apply_v093(in_path, out_path, template_path=TEMPLATE)

    wb = load_workbook(out_path)
    assert wb.sheetnames == SHEET_ORDER

    ws = wb["Product"]
    mandatory_row = [c.value for c in ws[1]]
    headers = [c.value for c in ws[2]]
    assert headers == [h for h, _ in PRODUCT_COLUMNS]
    assert mandatory_row[headers.index("Product SKU")] == "Mandatory ->"
    assert mandatory_row[headers.index("Pack Unit Box")] is None
    assert mandatory_row[headers.index("Transformation GHG")] == "Automatic"
    by_header = {c.value: c for c in ws[2]}
    assert by_header["Product SKU"].fill.fgColor.rgb == BLOCK_COLORS["Produit"]
    assert by_header["Supplier PCF value"].fill.fgColor.rgb == BLOCK_COLORS["Supplier PCF"]
    assert by_header["Transformation Process Name"].fill.fgColor.rgb == BLOCK_COLORS["Transformation"]
    assert by_header["Transformation DQR value"].fill.fgColor.rgb == "FFA6A6A6"
    assert by_header["Transformation GHG"].fill.fgColor.rgb == "FFA6A6A6"
    assert by_header["Data validation flag"].fill.fgColor.rgb == BLOCK_COLORS["Flag"]
    assert by_header["Product Details"].fill.fgColor.rgb == "FFD9D9D9"
    assert by_header["Supplier PCF"].fill.fgColor.rgb == "FFD9D9D9"
    assert ws.freeze_panes == "C3"

    grouped = {
        key
        for key, dim in ws.column_dimensions.items()
        if dim.outlineLevel == 1
    }
    first = headers.index("Category Code") + 1
    last = headers.index("Stock unit") + 1
    assert {get_column_letter(c) for c in range(first, last + 1)} <= grouped
    supp_first = headers.index("Supplier PCF framework") + 1
    supp_last = headers.index("Supplier PCF external review") + 1
    assert {get_column_letter(c) for c in range(supp_first, supp_last + 1)} <= grouped
    for header in ("Product Details", "Supplier PCF", "Transformation Details"):
        col = headers.index(header) + 1
        for r in (3, 202):
            assert ws.cell(row=r, column=col).fill.fgColor.rgb == "FFD9D9D9"

    assert ws["A3"].value == "SORHOY15"
    assert ws["H3"].value == 1.25
    energy_col = headers.index("Transformation Energy Type") + 1
    assert ws.cell(row=3, column=energy_col).value == "electricity"
    scrap_col = headers.index("Transformation Process Scrap Rate") + 1
    assert ws.cell(row=3, column=scrap_col).value == 0.03

    def data_cell(header: str):
        return ws.cell(row=3, column=headers.index(header) + 1)

    assert str(data_cell("Transformation DQR value").value).startswith("=IF(COUNT(")
    assert str(data_cell("Transformation GHG").value).startswith("=IF(")
    assert str(data_cell("Data validation flag").value).startswith("=IF(AND(")
    assert data_cell("Transformation DQR value").number_format == "0.000"
    assert data_cell("Supplier PDS").number_format == "0.00%"

    dvs = ws.data_validations.dataValidation
    formulas = {dv.formula1 for dv in dvs}
    assert '"PACT,TfS,ISO14067,other"' in formulas
    assert '"1,2,3,4,5"' in formulas

    dqr = wb[DQR_GUIDE_SHEET]
    values = {c.value for row in dqr.iter_rows() for c in row if c.value}
    assert "PACT Methodology 3.0 - Page 66" in values
    assert "Technological representativeness (TECH DQR)" in values
    assert "Excellent quality" in values

    ug = wb["UserGuide"]
    texts = [c.value for row in ug.iter_rows() for c in row if c.value]
    assert any("How to collect the data from suppliers" in str(t) for t in texts)
    assert any("Do you hold the PCF of [SKU]?" in str(t) for t in texts)


def test_tab_order_and_colors(tmp_path):
    in_path = _make_workbook(tmp_path)
    out_path = tmp_path / "out.xlsx"
    apply_v093(in_path, out_path, template_path=TEMPLATE)

    wb = load_workbook(out_path)
    assert wb.sheetnames == SHEET_ORDER
    for name, color in TAB_COLORS.items():
        ws = wb[name]
        assert ws.sheet_properties.tabColor is not None
        assert ws.sheet_properties.tabColor.rgb == color


def test_dqr_guide_title_fill(tmp_path):
    in_path = _make_workbook(tmp_path)
    out_path = tmp_path / "out.xlsx"
    apply_v093(in_path, out_path, template_path=TEMPLATE)

    wb = load_workbook(out_path)
    dqr = wb[DQR_GUIDE_SHEET]
    for coord in ("A4", "D4", "A12", "D12"):
        assert dqr[coord].fill.fgColor.rgb == "FF13501B"


def test_guides_format_preserved(tmp_path):
    in_path = _make_workbook(tmp_path)
    out_path = tmp_path / "out.xlsx"
    apply_v093(in_path, out_path, template_path=TEMPLATE)

    wb = load_workbook(out_path)
    ref = load_workbook(TEMPLATE)
    for name in ("UserGuide", "DQR_GUIDE_SHEET"):
        pass
    for name in ("UserGuide", DQR_GUIDE_SHEET):
        ws_out, ws_ref = wb[name], ref[name]
        assert [str(m) for m in ws_out.merged_cells.ranges] == [str(m) for m in ws_ref.merged_cells.ranges]
        for key, dim in ws_ref.column_dimensions.items():
            if dim.width:
                assert ws_out.column_dimensions[key].width == dim.width, (name, key)
        for key, dim in ws_ref.row_dimensions.items():
            if dim.height:
                assert ws_out.row_dimensions[key].height == dim.height, (name, key)
        for row_out, row_ref in zip(ws_out.iter_rows(), ws_ref.iter_rows()):
            for c_out, c_ref in zip(row_out, row_ref):
                assert c_out.value == c_ref.value
                if c_ref.fill.patternType == "solid":
                    assert c_out.fill.patternType == "solid"
                    assert c_out.fill.fgColor.rgb == c_ref.fill.fgColor.rgb
                if c_ref.font.b:
                    assert c_out.font.b


def test_real_file_if_available(tmp_path):
    if not INPUT.is_file():
        return
    out_path = tmp_path / "out.xlsx"
    apply_v093(INPUT, out_path, template_path=TEMPLATE)
    wb = load_workbook(out_path)
    assert wb.sheetnames == SHEET_ORDER
    comp = wb["Component"]
    by_header = {c.value: c for c in comp[2]}
    assert by_header["Component SKU"].fill.fgColor.rgb == BLOCK_COLORS["Composants"]
    assert by_header["Net Weight"].fill.fgColor.rgb == BLOCK_COLORS["Description du composant"]
    assert by_header["Supplier PDS"].fill.fgColor.rgb == BLOCK_COLORS["Supplier PCF"]
    assert by_header["Transformation Process Name"].fill.fgColor.rgb == BLOCK_COLORS["Transformation"]
    assert by_header["Component Details"].fill.fgColor.rgb == "FFD9D9D9"
    prod = wb["Product"]
    by_header_p = {c.value: c for c in prod[2]}
    assert by_header_p["Product SKU"].fill.fgColor.rgb == BLOCK_COLORS["Produit"]
    assert by_header_p["Product SKU"].font.color.rgb == "FFFFFFFF"
    comp2 = wb["Component"]
    by_header_c = {c.value: c for c in comp2[2]}
    assert by_header_c["UVP description"].fill.fgColor.rgb == BLOCK_COLORS["Description du composant"]


def test_prefill(tmp_path):
    spec = Path(__file__).resolve().parents[1] / "input" / "Spec v0.95.xlsx"
    material = (
        Path(__file__).resolve().parents[1]
        / "input"
        / "MaterialPackaging_ExtractPourPCF.xlsx"
    )
    if not spec.is_file() or not material.is_file():
        return
    in_path = _make_workbook(tmp_path)
    out_path = tmp_path / "out_prefilled.xlsx"
    counts = apply_v093(
        in_path,
        out_path,
        template_path=TEMPLATE,
        spec_path=spec,
        material_path=material,
    )
    assert counts is not None
    n_products, n_components = counts
    assert n_products == 217
    assert n_components > 900

    wb = load_workbook(out_path)
    prod = wb["Product"]
    headers = [c.value for c in prod[2]]
    sku_col = headers.index("Product SKU") + 1
    skus = [prod.cell(row=r, column=sku_col).value for r in range(3, 3 + n_products)]
    assert skus[0] is not None and all(s for s in skus)
    assert "SORHOY15" in skus
    comp = wb["Component"]
    headers_c = [c.value for c in comp[2]]
    csku_col = headers_c.index("Component SKU") + 1
    cskus = [comp.cell(row=r, column=csku_col).value for r in range(3, 3 + n_components)]
    assert cskus[0] is not None
    assert len(set(cskus)) == n_components
    # aucune ligne pré-remplie au-delà des données réelles
    assert prod.cell(row=3 + n_products, column=sku_col).value is None
    assert comp.cell(row=3 + n_components, column=csku_col).value is None
    # les formules par défaut restent en place sur les lignes pré-remplies
    dqr_col = headers.index("Transformation DQR value") + 1
    assert str(prod.cell(row=3, column=dqr_col).value).startswith("=IF(COUNT(")
    flag_col = headers.index("Data validation flag") + 1
    assert str(prod.cell(row=3, column=flag_col).value).startswith("=IF(AND(")
