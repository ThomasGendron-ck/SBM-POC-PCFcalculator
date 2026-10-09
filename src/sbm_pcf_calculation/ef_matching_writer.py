"""Génération du fichier EF matching (spec v0.98, bloc MissingEF_Matching).

Reprend le format du Fichier de collecte produit par ``datashare.py`` :
- en-têtes en ligne 2, ligne 1 = champ obligatoire (italique, rouge si Yes) ;
- couleurs de blocs par type de champ (automatique, Expert User, calculé,
  GROUPING, flag) ;
- colonnes GROUPING verticales grisées sur toute la hauteur des données ;
- formules Excel explicites pour les DQR value (moyenne GEO/TECH/TEMP) et le
  Data validation flag (règles PDS/DQR/FE manquant de la spec) ;
- listes déroulantes (Yes/No, DQR 1-5, impact High/Medium/Low) ;
- volets figés en C3, filtres automatiques, largeurs ajustées, couleur d'onglet.

Contenu pré-rempli : une ligne par composant sans FE du Fichier de collecte,
avec les matchings ecoinvent existants (``match_missing_fe``) reportés dans le
bloc RM AutoMatch, et les attributs composant (catégorie, fournisseur, pays,
matière, recyclé) issus de la collecte.
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .collect import compute_dqr
from .datashare import (
    DATA_END_ROW,
    DATA_START_ROW,
    GROUPING,
    HEADER_ROW,
    MANDATORY_ROW,
    _cell,
    _column_of,
    _dqr_flag,
    _pds_flag,
)
from .ecoinvent import supplier_country

EF_MATCHING_SHEET = "EF Matching"
EF_MATCHING_TAB_COLOR = "FF75A67C"
EF_SOURCE = "EcoInvent 3.12 cut-off (EF v3.1 GWP100)"

CALCULATED_FIELDS = {
    "RM AutoMatch DQR value",
    "Transfo AutoMatch DQR value",
    "RM UserValidation DQR value",
    "Transfo UserValidation DQR value",
}

BLOCK_COLORS = {
    "Composants": "FF75A67C",
    "Description du composant": "FFA3C4A7",
    "AutoMatch": "FF13501B",
    "Transfo AutoMatch": "FFF6C6AD",
    "UserValidation": "FFFBE3D6",
    "Transfo UserValidation": "FFF6C6AD",
    "Calculé": "FFA6A6A6",
    "Flag": "FFFFC000",
    GROUPING: "FFD9D9D9",
}
BLOCK_FONT_COLORS = {
    "Composants": "FF000000",
    "Description du composant": "FF000000",
    "AutoMatch": "FFFFFFFF",
    "Transfo AutoMatch": "FF000000",
    "UserValidation": "FF000000",
    "Transfo UserValidation": "FF000000",
    "Calculé": "FF000000",
    "Flag": "FF000000",
    GROUPING: "FF404040",
}

MANDATORY = {
    "Component SKU": "Mandatory ->",
    "Component Designation": "Yes",
    "Component Category Code": "Yes",
    "Component Category description": "Yes",
    "Component Carbon category": "Yes",
    "Component Carbon sub category": "Yes, if applicable",
    "Component Supplier Code": "Yes",
    "Component Supplier Name": "Yes",
    "Component Supplier Country": "Automatic",
    "Component Pack Unit box": "Yes",
    "UVP description": "Yes",
    "Raw Material (MB Product)": "Yes",
    "Raw Material - Carbon Footprint": "Yes, if available",
    "Recycled %": "Yes, if applicable",
    "RM AutoMatch EF Name": "Automatic",
    "RM AutoMatch EF Rationale": "Automatic",
    "RM AutoMatch EF Value": "Automatic",
    "RM AutoMatch EF Unit": "Automatic",
    "RM AutoMatch EF Geography": "Automatic",
    "RM AutoMatch EF Source": "Automatic",
    "RM AutoMatch EF PDS": "Automatic",
    "RM AutoMatch GEO DQR": "Automatic",
    "RM AutoMatch TECH DQR": "Automatic",
    "RM AutoMatch TEMP DQR": "Automatic",
    "RM AutoMatch DQR value": "Automatic",
    "Transfo AutoMatch Process Name": "Automatic",
    "Transfo AutoMatch EF Rationale": "Automatic",
    "Transfo AutoMatch Process EF Name": "Automatic",
    "Transfo AutoMatch Process EF Value": "Automatic",
    "Transfo AutoMatch Process EF Unit": "Automatic",
    "Transfo AutoMatch Process EF Source": "Automatic",
    "Transfo AutoMatch PDS value": "Automatic",
    "Transfo AutoMatch GEO DQR": "Automatic",
    "Transfo AutoMatch TECH DQR": "Automatic",
    "Transfo AutoMatch TEMP DQR": "Automatic",
    "Transfo AutoMatch DQR value": "Automatic",
    "RM UserValidation EF Yes/No": "Yes",
    "RM UserValidation EF Rationale": "Yes",
    "RM UserValidation EF Name": "Yes if EF not approved",
    "RM UserValidation EF Rationale (2)": "Yes if EF not approved",
    "RM UserValidation EF Value": "Yes if EF not approved",
    "RM UserValidation EF Unit": "Yes if EF not approved",
    "RM UserValidation EF Geography": "Yes if EF not approved",
    "RM UserValidation EF Source": "Yes if EF not approved",
    "RM UserValidation EF PDS": "Yes if EF not approved",
    "RM UserValidation GEO DQR": "Yes if EF not approved",
    "RM UserValidation TECH DQR": "Yes if EF not approved",
    "RM UserValidation TEMP DQR": "Yes if EF not approved",
    "RM UserValidation DQR value": "Automatic",
    "Transfo UserValidation Process Name": "Yes, if User validation",
    "Transfo UserValidation Process EF Name": "Yes, if User validation",
    "Transfo UserValidation Process EF Value": "Yes, if User validation",
    "Transfo UserValidation Process EF Unit": "Yes, if User validation",
    "Transfo UserValidation Process EF Source": "Yes, if User validation",
    "Transfo UserValidation PDS value": "Yes, if User validation",
    "Transfo UserValidation GEO DQR": "Yes, if User validation",
    "Transfo UserValidation TECH DQR": "Yes, if User validation",
    "Transfo UserValidation TEMP DQR": "Yes, if User validation",
    "Transfo UserValidation DQR value": "Automatic",
    "Data validation flag": "Automatic",
    "Data validation flag impact": "Automatic",
}

# Renommage Excel : la spec v0.98 contient deux colonnes « RM UserValidation EF
# Rationale » ; la seconde est suffixée en interne puis renommée à l'écriture.
EXCEL_HEADER_MAP = {"RM UserValidation EF Rationale (2)": "RM UserValidation EF Rationale"}

# Colonnes du fichier (ordre exact du bloc MissingEF_Matching, spec v0.98,
# lignes 128-192) : (en-tête interne, bloc).
EF_MATCHING_COLUMNS: list[tuple[str, str]] = [
    ("Component SKU", "Composants"),
    ("Component Designation", "Composants"),
    ("Component Category Code", "Composants"),
    ("Component Category description", "Composants"),
    ("Component Carbon category", "Composants"),
    ("Component Carbon sub category", "Composants"),
    ("Component Supplier Code", "Composants"),
    ("Component Supplier Name", "Composants"),
    ("Component Supplier Country", "Composants"),
    ("Component Pack Unit box", "Composants"),
    ("Component Details", GROUPING),
    ("UVP description", "Description du composant"),
    ("Raw Material (MB Product)", "Description du composant"),
    ("Raw Material - Carbon Footprint", "Description du composant"),
    ("Recycled %", "Description du composant"),
    ("Component Material", GROUPING),
    ("RM AutoMatch EF Name", "AutoMatch"),
    ("RM AutoMatch EF Rationale", "AutoMatch"),
    ("RM AutoMatch EF Value", "AutoMatch"),
    ("RM AutoMatch EF Unit", "AutoMatch"),
    ("RM AutoMatch EF Geography", "AutoMatch"),
    ("RM AutoMatch EF Source", "AutoMatch"),
    ("RM AutoMatch EF PDS", "AutoMatch"),
    ("RM AutoMatch GEO DQR", "AutoMatch"),
    ("RM AutoMatch TECH DQR", "AutoMatch"),
    ("RM AutoMatch TEMP DQR", "AutoMatch"),
    ("RM AutoMatch DQR value", "Calculé"),
    ("AutoMatch Details", GROUPING),
    ("Transfo AutoMatch Process Name", "Transfo AutoMatch"),
    ("Transfo AutoMatch EF Rationale", "Transfo AutoMatch"),
    ("Transfo AutoMatch Process EF Name", "Transfo AutoMatch"),
    ("Transfo AutoMatch Process EF Value", "Transfo AutoMatch"),
    ("Transfo AutoMatch Process EF Unit", "Transfo AutoMatch"),
    ("Transfo AutoMatch Process EF Source", "Transfo AutoMatch"),
    ("Transfo AutoMatch PDS value", "Transfo AutoMatch"),
    ("Transfo AutoMatch GEO DQR", "Transfo AutoMatch"),
    ("Transfo AutoMatch TECH DQR", "Transfo AutoMatch"),
    ("Transfo AutoMatch TEMP DQR", "Transfo AutoMatch"),
    ("Transfo AutoMatch DQR value", "Calculé"),
    ("Transfo AutoMatch Details", GROUPING),
    ("RM UserValidation EF Yes/No", "UserValidation"),
    ("RM UserValidation EF Rationale", "UserValidation"),
    ("RM UserValidation EF Name", "UserValidation"),
    ("RM UserValidation EF Rationale (2)", "UserValidation"),
    ("RM UserValidation EF Value", "UserValidation"),
    ("RM UserValidation EF Unit", "UserValidation"),
    ("RM UserValidation EF Geography", "UserValidation"),
    ("RM UserValidation EF Source", "UserValidation"),
    ("RM UserValidation EF PDS", "UserValidation"),
    ("RM UserValidation GEO DQR", "UserValidation"),
    ("RM UserValidation TECH DQR", "UserValidation"),
    ("RM UserValidation TEMP DQR", "UserValidation"),
    ("RM UserValidation DQR value", "Calculé"),
    ("UserValidation", GROUPING),
    ("Transfo UserValidation Process Name", "Transfo UserValidation"),
    ("Transfo UserValidation Process EF Name", "Transfo UserValidation"),
    ("Transfo UserValidation Process EF Value", "Transfo UserValidation"),
    ("Transfo UserValidation Process EF Unit", "Transfo UserValidation"),
    ("Transfo UserValidation Process EF Source", "Transfo UserValidation"),
    ("Transfo UserValidation PDS value", "Transfo UserValidation"),
    ("Transfo UserValidation GEO DQR", "Transfo UserValidation"),
    ("Transfo UserValidation TECH DQR", "Transfo UserValidation"),
    ("Transfo UserValidation TEMP DQR", "Transfo UserValidation"),
    ("Transfo UserValidation DQR value", "Calculé"),
    ("Transfo UserValidation Details", GROUPING),
    ("Data validation flag", "Flag"),
    ("Data validation flag impact", "Flag"),
]

# Groupings Plan : (première, dernière) colonne du groupe, bornes inclusives.
GROUPING_RANGES = [
    ("UVP description", "Recycled %"),
    ("RM AutoMatch EF Name", "RM AutoMatch TEMP DQR"),
    ("Transfo AutoMatch Process Name", "Transfo AutoMatch TEMP DQR"),
    ("RM UserValidation EF Yes/No", "RM UserValidation TEMP DQR"),
    ("Transfo UserValidation Process Name", "Transfo UserValidation TEMP DQR"),
]

NUMBER_FORMATS = {
    "Component Pack Unit box": "0.000",
    "Recycled %": "0.00%",
    "RM AutoMatch EF Value": "0.00000",
    "RM AutoMatch EF PDS": "0.00%",
    "RM AutoMatch GEO DQR": "0.000",
    "RM AutoMatch TECH DQR": "0.000",
    "RM AutoMatch TEMP DQR": "0.000",
    "RM AutoMatch DQR value": "0.000",
    "Transfo AutoMatch Process EF Value": "0.00000",
    "Transfo AutoMatch PDS value": "0.00%",
    "Transfo AutoMatch GEO DQR": "0.000",
    "Transfo AutoMatch TECH DQR": "0.000",
    "Transfo AutoMatch TEMP DQR": "0.000",
    "Transfo AutoMatch DQR value": "0.000",
    "RM UserValidation EF Value": "0.00000",
    "RM UserValidation EF PDS": "0.00%",
    "RM UserValidation GEO DQR": "0.000",
    "RM UserValidation TECH DQR": "0.000",
    "RM UserValidation TEMP DQR": "0.000",
    "RM UserValidation DQR value": "0.000",
    "Transfo UserValidation Process EF Value": "0.00000",
    "Transfo UserValidation PDS value": "0.00%",
    "Transfo UserValidation GEO DQR": "0.000",
    "Transfo UserValidation TECH DQR": "0.000",
    "Transfo UserValidation TEMP DQR": "0.000",
    "Transfo UserValidation DQR value": "0.000",
}

LIST_RULES = {
    "RM UserValidation EF Yes/No": '"Yes,No"',
    "Data validation flag impact": '"High,Medium,Low,No impact"',
}

# Bloc de chaque en-tête (pour les modules de test / réutilisation).
HEADER_BLOCKS = {header: block for header, block in EF_MATCHING_COLUMNS}


def _dqr_formula(prefix: str, row: int) -> str:
    geo = _cell(EF_MATCHING_COLUMNS, f"{prefix} GEO DQR", row)
    tech = _cell(EF_MATCHING_COLUMNS, f"{prefix} TECH DQR", row)
    temp = _cell(EF_MATCHING_COLUMNS, f"{prefix} TEMP DQR", row)
    return f'=IF(COUNT({geo},{tech},{temp})=0,"",ROUND(AVERAGE({geo},{tech},{temp}),3))'


def _flag_formula(row: int) -> str:
    def c(header: str) -> str:
        return _cell(EF_MATCHING_COLUMNS, header, row)

    sku = c("Component SKU")
    parts = [
        f'IF(AND({sku}<>"",{c("Component Supplier Code")}=""),"Supplier code empty (High);","")',
        (
            f'IF(AND({sku}<>"",{c("UVP description")}="",'
            f'{c("Raw Material (MB Product)")}="",'
            f'{c("Raw Material - Carbon Footprint")}=""),'
            f'"No raw material available (High);","")'
        ),
        (
            f'IF(AND({sku}<>"",{c("RM AutoMatch EF Name")}="",'
            f'{c("RM UserValidation EF Name")}=""),'
            f'"No EF associated to the material (High);","")'
        ),
        _pds_flag(c("RM AutoMatch EF PDS")),
        _dqr_flag(c("RM AutoMatch DQR value")),
        _pds_flag(c("Transfo AutoMatch PDS value")),
        _dqr_flag(c("Transfo AutoMatch DQR value")),
        (
            f'IF(AND({c("RM UserValidation EF Yes/No")}="Yes",'
            f'{c("RM UserValidation EF Name")}=""),'
            f'"User validation without EF name (High);","")'
        ),
        _pds_flag(c("RM UserValidation EF PDS")),
        _dqr_flag(c("RM UserValidation DQR value")),
        _pds_flag(c("Transfo UserValidation PDS value")),
        _dqr_flag(c("Transfo UserValidation DQR value")),
    ]
    return "=" + "&".join(parts) + '&""'


def _defaults(row: int) -> dict[str, object]:
    return {
        "RM AutoMatch EF PDS": 0,
        "Transfo AutoMatch PDS value": 0,
        "RM UserValidation EF PDS": 0,
        "Transfo UserValidation PDS value": 0,
        "RM AutoMatch DQR value": _dqr_formula("RM AutoMatch", row),
        "Transfo AutoMatch DQR value": _dqr_formula("Transfo AutoMatch", row),
        "RM UserValidation DQR value": _dqr_formula("RM UserValidation", row),
        "Transfo UserValidation DQR value": _dqr_formula("Transfo UserValidation", row),
        "Data validation flag": _flag_formula(row),
    }


def build_ef_matching_rows(collecte: pd.DataFrame, matching: pd.DataFrame | None) -> list[dict]:
    """Une ligne par composant sans FE, pré-remplie depuis la collecte et le matching ecoinvent."""
    no_fe = collecte[collecte["RM EF Value"].isna() & collecte["Component SKU"].notna()].copy()
    if no_fe.empty:
        return []
    match_by_sku: dict[str, pd.Series] = {}
    if matching is not None and not matching.empty:
        match_by_sku = {str(r["Component SKU"]): r for _, r in matching.iterrows()}
    def _first_valid(lines: pd.DataFrame, col: str):
        if col not in lines.columns:
            return None
        vals = lines[col].dropna()
        vals = vals[vals.astype(str).str.strip() != ""]
        return vals.iloc[0] if not vals.empty else None

    rows: list[dict] = []
    for sku, lines in no_fe.groupby("Component SKU"):
        first = lines.iloc[0]
        row = {
            "Component SKU": sku,
            "Component Designation": _first_valid(lines, "Component Designation") or first["Component Designation"],
            "Component Category Code": _first_valid(lines, "Category Code"),
            "Component Category description": _first_valid(lines, "Category description"),
            "Component Carbon category": _first_valid(lines, "Carbon category"),
            "Component Supplier Code": _first_valid(lines, "Supplier code"),
            "Component Supplier Name": _first_valid(lines, "Supplier Name"),
            "Component Supplier Country": supplier_country(_first_valid(lines, "Supplier code")),
            "Component Pack Unit box": _first_valid(lines, "Pack unit box"),
            "UVP description": _first_valid(lines, "UVP description"),
            "Raw Material (MB Product)": _first_valid(lines, "Raw Material"),
            "Raw Material - Carbon Footprint": _first_valid(lines, "Raw Material - Carbon Footprint"),
            "Recycled %": _first_valid(lines, "Recycled %"),
        }
        m = match_by_sku.get(str(sku))
        if m is not None and m["Statut"] == "MATCHÉ" and pd.notna(m["FE proposé (kg CO2e/kg)"]):
            geo = m["Géographie"]
            dqr = compute_dqr(geo, EF_SOURCE, True)
            row.update(
                {
                    "RM AutoMatch EF Name": m["Dataset ecoinvent"],
                    "RM AutoMatch EF Rationale": m["Règle de matching"],
                    "RM AutoMatch EF Value": float(m["FE proposé (kg CO2e/kg)"]),
                    "RM AutoMatch EF Unit": "kgCO2e/kg",
                    "RM AutoMatch EF Geography": geo,
                    "RM AutoMatch EF Source": EF_SOURCE,
                    "RM AutoMatch GEO DQR": dqr["RM GEO DQR"],
                    "RM AutoMatch TECH DQR": dqr["RM TECH DQR"],
                    "RM AutoMatch TEMP DQR": dqr["RM TEMP DQR"],
                }
            )
        rows.append(row)
    return rows


def _write_rows(ws, rows: list[dict]) -> None:
    for offset, row in enumerate(rows, start=DATA_START_ROW):
        values = {**_defaults(offset), **row}
        for header, value in values.items():
            if value is None:
                continue
            try:
                col = _column_of(EF_MATCHING_COLUMNS, header)
            except KeyError:
                continue
            ws.cell(row=offset, column=col, value=value)


def _header_style(cell, block: str) -> None:
    cell.fill = PatternFill("solid", fgColor=BLOCK_COLORS[block])
    cell.font = Font(color=BLOCK_FONT_COLORS[block], bold=True)
    if block == GROUPING:
        cell.alignment = Alignment(vertical="center", textRotation=90, wrap_text=True)
    else:
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _build_sheet(wb, rows: list[dict]):
    if EF_MATCHING_SHEET in wb.sheetnames:
        del wb[EF_MATCHING_SHEET]
    ws = wb.create_sheet(EF_MATCHING_SHEET)
    data_end_row = max(DATA_START_ROW + len(rows) - 1, DATA_START_ROW)
    for idx, (header, block) in enumerate(EF_MATCHING_COLUMNS, start=1):
        if header in CALCULATED_FIELDS:
            block = "Calculé"
        excel_header = EXCEL_HEADER_MAP.get(header, header)
        cell = ws.cell(row=HEADER_ROW, column=idx, value=excel_header)
        _header_style(cell, block)
        mandatory = MANDATORY.get(header)
        mcell = ws.cell(row=MANDATORY_ROW, column=idx, value=mandatory)
        mcell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        mcell.font = Font(italic=True, color="FF404040")
        if str(mandatory or "").startswith(("Yes", "Mandatory")):
            mcell.font = Font(italic=True, bold=True, color="FFC00000")
        letter = get_column_letter(idx)
        if block == GROUPING:
            ws.column_dimensions[letter].width = 3.43
        elif header == "Data validation flag":
            ws.column_dimensions[letter].width = 45
        elif header == "Data validation flag impact":
            ws.column_dimensions[letter].width = 15
        else:
            ws.column_dimensions[letter].width = 18
    ws.row_dimensions[MANDATORY_ROW].height = 30
    ws.row_dimensions[HEADER_ROW].height = 60
    _write_rows(ws, rows)
    for header, fmt in NUMBER_FORMATS.items():
        col = _column_of(EF_MATCHING_COLUMNS, header)
        for row_idx in range(DATA_START_ROW, data_end_row + 1):
            ws.cell(row=row_idx, column=col).number_format = fmt
    for header, formula in LIST_RULES.items():
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        col = _column_of(EF_MATCHING_COLUMNS, header)
        dv.add(f"{get_column_letter(col)}{DATA_START_ROW}:{get_column_letter(col)}{data_end_row}")
        ws.add_data_validation(dv)
    dqr_dv = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    for header in (
        "RM AutoMatch GEO DQR",
        "RM AutoMatch TECH DQR",
        "RM AutoMatch TEMP DQR",
        "Transfo AutoMatch GEO DQR",
        "Transfo AutoMatch TECH DQR",
        "Transfo AutoMatch TEMP DQR",
        "RM UserValidation GEO DQR",
        "RM UserValidation TECH DQR",
        "RM UserValidation TEMP DQR",
        "Transfo UserValidation GEO DQR",
        "Transfo UserValidation TECH DQR",
        "Transfo UserValidation TEMP DQR",
    ):
        col = _column_of(EF_MATCHING_COLUMNS, header)
        dqr_dv.add(f"{get_column_letter(col)}{DATA_START_ROW}:{get_column_letter(col)}{data_end_row}")
    ws.add_data_validation(dqr_dv)
    grouping_fill = PatternFill("solid", fgColor=BLOCK_COLORS[GROUPING])
    for idx, (_, block) in enumerate(EF_MATCHING_COLUMNS, start=1):
        if block != GROUPING:
            continue
        for row_idx in range(DATA_START_ROW, data_end_row + 1):
            ws.cell(row=row_idx, column=idx).fill = grouping_fill
    for first, last in GROUPING_RANGES:
        first_col = _column_of(EF_MATCHING_COLUMNS, first)
        last_col = _column_of(EF_MATCHING_COLUMNS, last)
        for col in range(first_col, last_col + 1):
            dim = ws.column_dimensions[get_column_letter(col)]
            dim.outlineLevel = 1
            dim.hidden = False
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(EF_MATCHING_COLUMNS))}{data_end_row}"
    for idx, (header, block) in enumerate(EF_MATCHING_COLUMNS, start=1):
        if block == GROUPING:
            continue
        longest = len(header)
        for row_idx in range(DATA_START_ROW, data_end_row + 1):
            value = ws.cell(row=row_idx, column=idx).value
            if value is not None and not str(value).startswith("="):
                longest = max(longest, len(str(value)))
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, 40)
    ws.sheet_properties.tabColor = EF_MATCHING_TAB_COLOR
    return ws


def write_ef_matching(
    output_path: str | Path,
    collecte: pd.DataFrame,
    matching: pd.DataFrame | None = None,
    template_path: str | Path | None = None,
) -> int:
    """Génère le fichier EF matching en un seul onglet (EF Matching).

    Retourne le nombre de composants sans FE écrits.
    """
    rows = build_ef_matching_rows(collecte, matching)
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    _build_sheet(wb, rows)
    wb.save(output_path)
    return len(rows)
