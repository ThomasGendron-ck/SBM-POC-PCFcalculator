"""Génération du fichier de saisie des données de transformation (à remplir par SBM)."""

from pathlib import Path

import pandas as pd

from . import config


def build_saisie_transformation(collecte: pd.DataFrame) -> pd.DataFrame:
    """Une ligne par couple (produit, composant) unique, pré-remplie avec le fournisseur."""
    saisie = (
        collecte[["Product SKU", "Component SKU", "Supplier Name"]]
        .dropna(subset=["Component SKU"])
        .drop_duplicates(subset=["Product SKU", "Component SKU"])
        .reset_index(drop=True)
    )
    for col in config.SAISIE_COLUMNS:
        if col not in saisie.columns:
            saisie[col] = None
    return saisie[config.SAISIE_COLUMNS]


def write_saisie_transformation(collecte: pd.DataFrame, output_path: str) -> None:
    """Écrit le classeur de saisie avec un onglet Guide et un onglet Saisie transformation."""
    saisie = build_saisie_transformation(collecte)
    guide = pd.DataFrame(config.SAISIE_GUIDE_ROWS, columns=["Colonne", "Explanation"])
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        guide.to_excel(writer, sheet_name=config.SAISIE_GUIDE_SHEET, index=False)
        saisie.to_excel(writer, sheet_name=config.SAISIE_TRANSFORMATION_SHEET, index=False)
    _style_saisie(output_path)


def _style_saisie(output_path: str) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = load_workbook(output_path)
    guide = wb[config.SAISIE_GUIDE_SHEET]
    guide.column_dimensions["A"].width = 34
    guide.column_dimensions["B"].width = 110
    header_fill = PatternFill("solid", fgColor="FF13501B")
    for cell in guide[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    for row in guide.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    saisie_ws = wb[config.SAISIE_TRANSFORMATION_SHEET]
    widths = [16, 16, 30, 34, 26, 22, 16, 30]
    for i, width in enumerate(widths, start=1):
        saisie_ws.column_dimensions[saisie_ws.cell(row=1, column=i).column_letter].width = width
    prefilled = {"A", "B", "C"}
    pre_fill = PatternFill("solid", fgColor="FFE7E6E6")
    for cell in saisie_ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFFFF", bold=True)
    for row_idx in range(2, saisie_ws.max_row + 1):
        for col_letter in prefilled:
            saisie_ws[f"{col_letter}{row_idx}"].fill = pre_fill
    saisie_ws.freeze_panes = "A2"
    wb.save(output_path)
