"""Génération des rapports : statistiques sur les facteurs d'émission manquants et priorisation."""

import pandas as pd

from . import config


def collecte_stats(collecte: pd.DataFrame) -> pd.DataFrame:
    """Statistiques de complétude du Fichier de collecte (produits, FE manquants, flags)."""
    pcf = collecte.groupby("Product SKU")["PCF Value"].first()
    sans_fe = collecte[collecte["Flag"].fillna("").str.contains("PAS DE FE")]
    return pd.DataFrame(
        [
            {"Indicateur": "Produits", "Valeur": collecte["Product SKU"].nunique()},
            {"Indicateur": "Lignes composants", "Valeur": len(collecte)},
            {"Indicateur": "Produits avec PCF calculé", "Valeur": int(pcf.notna().sum())},
            {"Indicateur": "Produits sans PCF complet", "Valeur": int(pcf.isna().sum())},
            {"Indicateur": "Lignes sans FE", "Valeur": len(sans_fe)},
            {"Indicateur": "Composants uniques sans FE", "Valeur": sans_fe["Component SKU"].nunique()},
        ]
    )


def collecte_flags(collecte: pd.DataFrame) -> pd.DataFrame:
    """Distribution des flags de validation par ligne, avec description."""
    exploded = collecte.assign(Flag=collecte["Flag"].fillna("").str.split(" | ", regex=False)).explode("Flag")
    exploded = exploded[exploded["Flag"] != ""]
    flags = (
        exploded.groupby("Flag")
        .agg(**{"Nb lignes": ("Flag", "size"), "Nb produits": ("Product SKU", "nunique")})
        .sort_values("Nb lignes", ascending=False)
        .reset_index()
    )
    flags["Description"] = flags["Flag"].map(config.FLAG_DESCRIPTIONS).fillna("")
    return flags


def collecte_matieres_manquantes(collecte: pd.DataFrame) -> pd.DataFrame:
    """Composants uniques sans FE, groupés par matière, pour prioriser les recherches."""
    sans_fe = collecte[collecte["Flag"].fillna("").str.contains("PAS DE FE")].copy()
    if sans_fe.empty:
        return pd.DataFrame(columns=["Raw Material", "Nb composants uniques", "Nb produits impactés"])
    return (
        sans_fe.groupby("Raw Material", dropna=False)
        .agg(**{"Nb composants uniques": ("Component SKU", "nunique"), "Nb produits impactés": ("Product SKU", "nunique")})
        .sort_values("Nb produits impactés", ascending=False)
        .reset_index()
    )


def collecte_composants_prioritaires(collecte: pd.DataFrame) -> pd.DataFrame:
    """Composants uniques sans FE triés par nombre de produits impactés."""
    sans_fe = collecte[collecte["Flag"].fillna("").str.contains("PAS DE FE")].copy()
    if sans_fe.empty:
        return pd.DataFrame(columns=["Component SKU", "Component Designation", "Category Code", "Raw Material", "Nb produits impactés"])
    return (
        sans_fe.groupby(["Component SKU", "Component Designation", "Category Code", "Raw Material"], dropna=False)
        .agg(**{"Nb produits impactés": ("Product SKU", "nunique")})
        .sort_values("Nb produits impactés", ascending=False)
        .reset_index()
    )


MARKER_TEXT_ROTATION = 90

def _autosize_columns(ws, collecte: pd.DataFrame, header_row: int) -> None:
    """Ajuste la largeur des colonnes au contenu (en-tête et données)."""
    from openpyxl.utils import get_column_letter

    marker_index = {
        col for col in config.GROUPING_MARKER_COLUMNS if col in collecte.columns
    }
    for col_idx, col in enumerate(collecte.columns, start=1):
        letter = get_column_letter(col_idx)
        if col in marker_index:
            ws.column_dimensions[letter].width = config.GROUPING_MARKER_WIDTH
            continue
        widths = [len(str(col))]
        for value in collecte[col].head(200):
            if value is None or (isinstance(value, float) and pd.isna(value)):
                continue
            widths.append(len(str(value)))
        cap = None if col == "Flag" else 40
        width = max(widths) + 2
        ws.column_dimensions[letter].width = min(width, cap) if cap else width

def _style_collecte_sheet(ws, collecte: pd.DataFrame) -> None:
    """Applique les couleurs de blocs, colonnes marqueurs GROUPING verticales,
    formats nombre, groupements (Plans > Grouper), auto-ajustement des largeurs,
    filtres automatiques et volets figés du livrable (règles v0.7)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_row = config.COLLECTE_HEADER_ROW
    n_rows = len(collecte)
    index = {col: i + 1 for i, col in enumerate(collecte.columns)}
    marker_index = {col: idx for idx, col in enumerate(collecte.columns, start=1) if col in config.GROUPING_MARKER_COLUMNS}

    header_fills = {}
    header_fonts = {}
    current_color = None
    current_font = None
    for block_name, block_cols in config.BLOCK_HEADERS:
        if block_name is None:
            continue
        if block_name in config.BLOCK_COLORS:
            current_color = config.BLOCK_COLORS[block_name]
        if block_name in config.BLOCK_FONT_COLORS:
            current_font = config.BLOCK_FONT_COLORS[block_name]
        elif current_color is None:
            continue
        for col in block_cols:
            if col in index:
                header_fills[index[col]] = current_color
                header_fonts[index[col]] = current_font or "FF000000"

    for col_idx, color in header_fills.items():
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(color=header_fonts.get(col_idx, "FF000000"), bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = config.COLLECTE_HEADER_ROW_HEIGHT

    for col in marker_index:
        cell = ws.cell(row=header_row, column=marker_index[col])
        cell.alignment = Alignment(
            text_rotation=MARKER_TEXT_ROTATION,
            vertical="bottom",
            horizontal="center",
            wrap_text=True,
        )
        color = header_fills.get(marker_index[col])
        if color:
            for row_idx in range(header_row + 1, header_row + 1 + n_rows):
                ws.cell(row=row_idx, column=marker_index[col]).fill = PatternFill("solid", fgColor=color)

    flag_idx = index.get("Flag")
    if flag_idx:
        flag_fill = PatternFill("solid", fgColor=config.FLAG_FILL)
        for row_idx in range(header_row + 1, header_row + 1 + n_rows):
            ws.cell(row=row_idx, column=flag_idx).fill = flag_fill

    for col, number_format in config.NUMBER_FORMATS.items():
        if col not in index:
            continue
        for row_idx in range(header_row + 1, header_row + 1 + n_rows):
            ws.cell(row=row_idx, column=index[col]).number_format = number_format

    for first, last in config.COLUMN_GROUPS:
        if first not in index or last not in index:
            continue
        for col_idx in range(index[first], index[last] + 1):
            ws.column_dimensions[get_column_letter(col_idx)].outlineLevel = 1
    ws.sheet_properties.outlinePr.summaryRight = True
    first, last = config.COLLECTE_TOP_GROUP_ROWS
    for row_idx in range(first, last + 1):
        ws.row_dimensions[row_idx].outlineLevel = 1
        ws.row_dimensions[row_idx].hidden = True
    _autosize_columns(ws, collecte, header_row)
    ws.freeze_panes = config.COLLECTE_FREEZE_PANES
    last_col = get_column_letter(len(collecte.columns))
    ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row + n_rows}"


def _pds_ef_phase(collecte: pd.DataFrame, phase: str, carbon_category: str | None) -> pd.Series:
    """PDS EF d'une phase : part des lignes avec un FE renseigné (0 si aucun FE)."""
    if carbon_category is not None:
        mask = collecte["Carbon category"] == carbon_category
    elif phase == "Transformation":
        mask = collecte["Transformation GHG"].notna()
    else:
        mask = collecte["Freight GHG"].notna()
    lines = collecte[mask]
    if lines.empty:
        return pd.Series(0.0, index=collecte.index)
    has_ef = lines["RM EF Value"].notna() & (lines["RM EF Value"] > 0)
    return pd.Series(float(has_ef.mean()), index=collecte.index).where(mask, 0.0)


def build_synthese(collecte: pd.DataFrame) -> pd.DataFrame:
    """Onglet Synthèse : une ligne par produit avec le PCF/DQR/PDS total et la
    décomposition par phase (Raw Material, Packaging, Transformation, Freight,
    Use, End of Life). Use et End of Life ne sont pas encore alimentées
    (périmètre cradle-to-gate) : leurs colonnes existent mais restent vides.
    Pour chaque phase : PDS Activity Data, PDS EF, PDS Total, DQR, GHG Value et
    % PCF total, avec un groupement Excel entre PDS Activity Data et DQR.
    """
    rows = []
    for sku, lines in collecte.groupby("Product SKU"):
        pcf_total = lines["PCF Value"].dropna().iloc[0] if lines["PCF Value"].notna().any() else float("nan")
        dqr_total = lines["DQR Product"].dropna().iloc[0] if lines["DQR Product"].notna().any() else float("nan")
        pds_total = lines["PDS Product"].dropna().iloc[0] if lines["PDS Product"].notna().any() else float("nan")
        row = {
            "Product SKU": sku,
            "Product Designation": lines["Product Designation"].iloc[0],
            "PCF Value": pcf_total,
            "DQR": dqr_total,
            "PDS": pds_total,
        }
        phase_mask = {
            "Raw Material": lines["Carbon category"] == "Raw Material",
            "Packaging": lines["Carbon category"] == "PACKAGING",
            "Transformation": lines["Transformation GHG"].notna(),
            "Freight": lines["Freight GHG"].notna(),
            "Use": pd.Series(False, index=lines.index),
            "End of Life": pd.Series(False, index=lines.index),
        }
        for phase in ("Raw Material", "Packaging", "Transformation", "Freight", "Use", "End of Life"):
            sub = lines[phase_mask[phase]]
            ghg = float(sub["RM GHG"].sum(min_count=1))
            if phase == "Transformation":
                ghg = float(sub["Transformation GHG"].sum(min_count=1))
            elif phase == "Freight":
                ghg = float(sub["Freight GHG"].sum(min_count=1))
            pds_activity = float(sub["RM PDS Activity Data"].mean()) if len(sub) else float("nan")
            pds_ef = float(sub["RM EF PDS"].mean()) if len(sub) else float("nan")
            pds_phase = pds_activity * pds_ef if pd.notna(pds_activity) and pd.notna(pds_ef) else float("nan")
            dqr = float(sub["RM DQR value"].mean()) if len(sub) and sub["RM DQR value"].notna().any() else float("nan")
            row[f"{phase} - PDS Activity Data"] = pds_activity
            row[f"{phase} - PDS EF"] = pds_ef
            row[f"{phase} - PDS Total"] = pds_phase
            row[f"{phase} - DQR"] = dqr
            row[f"{phase} - GHG Value"] = ghg if pd.notna(ghg) else float("nan")
            row[f"{phase} - % PCF total"] = (
                ghg / pcf_total if pd.notna(ghg) and pd.notna(pcf_total) and pcf_total else float("nan")
            )
        rows.append(row)
    return pd.DataFrame(rows)


SYNTHESE_COLUMNS = ["Product SKU", "Product Designation", "PCF Value", "DQR", "PDS"] + [
    f"{phase} - {metric}"
    for phase in ("Raw Material", "Packaging", "Transformation", "Freight", "Use", "End of Life")
    for metric in ("PDS Activity Data", "PDS EF", "PDS Total", "DQR", "GHG Value", "% PCF total")
]


SYNTHESE_HEADER_HEIGHT = 60.0

def _style_synthese_sheet(ws, synthese: pd.DataFrame) -> None:
    """Style de l'onglet Synthèse : en-tête avec renvoi à la ligne, formats
    décimal(10,3) avec séparateur de milliers et % en décimal(3,2), groupe
    Excel PDS Activity Data -> DQR par phase."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="FF156082")
    phase_colors = {
        "Raw Material": "FFA3C4A7",
        "Packaging": "FFD1E1D3",
        "Transformation": "FFD1E1D3",
        "Freight": "FF6FC5E6",
        "Use": "FFF2AA84",
        "End of Life": "FFC7C7C7",
    }
    for col_idx, col in enumerate(synthese.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        phase = col.split(" - ")[0] if " - " in col else None
        if col in ("Product SKU", "Product Designation", "PCF Value", "DQR", "PDS"):
            cell.fill = header_fill
            cell.font = Font(color="FFFFFFFF", bold=True)
        elif phase in phase_colors:
            cell.fill = PatternFill("solid", fgColor=phase_colors[phase])
            cell.font = Font(color="FF000000", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        number_format = None
        if col.endswith("% PCF total"):
            number_format = "0.00%"
        elif col not in ("Product SKU", "Product Designation"):
            number_format = "#,##0.000"
        if number_format:
            for row_idx in range(2, len(synthese) + 2):
                ws.cell(row=row_idx, column=col_idx).number_format = number_format
        widths = [min(len(str(col)), 18)]
        for value in synthese[col].head(200):
            if value is None or (isinstance(value, float) and pd.isna(value)):
                continue
            widths.append(len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(widths) + 2, 40)
    for phase in ("Raw Material", "Packaging", "Transformation", "Freight", "Use", "End of Life"):
        first = f"{phase} - PDS Activity Data"
        last = f"{phase} - DQR"
        if first in synthese.columns and last in synthese.columns:
            for col_idx in range(synthese.columns.get_loc(first) + 1, synthese.columns.get_loc(last) + 2):
                ws.column_dimensions[get_column_letter(col_idx)].outlineLevel = 1
    ws.sheet_properties.outlinePr.summaryRight = True
    ws.row_dimensions[1].height = SYNTHESE_HEADER_HEIGHT
    ws.freeze_panes = "C2"


def write_collecte_report(collecte: pd.DataFrame, output_path: str, matching: pd.DataFrame | None = None) -> None:
    """
    Écrit le classeur Fichier de collecte avec onglets stats, priorisation et matching ecoinvent.
    L'onglet principal démarre en A6 (en-têtes en ligne 6, données à partir de la
    ligne 7), reprend les couleurs de blocs de la Clé de détermination v0.7, les
    formats nombre de la colonne « Type de données », les groupements de colonnes
    (Plans > Grouper), les filtres automatiques en ligne 6 et les volets figés en C7.
    Les en-têtes Excel sont les noms de la colonne B de la Clé (les colonnes
    internes du bloc produit sont préfixées « Product » puis renommées à l'écriture,
    doublons inclus comme dans la Clé).
    """
    excel_collecte = collecte.rename(columns=config.EXCEL_HEADER_MAP)
    synthese = build_synthese(collecte)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        synthese.to_excel(writer, sheet_name="Synthese", index=False)
        excel_collecte.to_excel(
            writer,
            sheet_name="Fichier de collecte",
            index=False,
            startrow=config.COLLECTE_HEADER_ROW - 1,
        )
        collecte_stats(collecte).to_excel(writer, sheet_name="Stats", index=False)
        collecte_flags(collecte).to_excel(writer, sheet_name="Flags", index=False)
        collecte_matieres_manquantes(collecte).to_excel(writer, sheet_name="FE manquants par matiere", index=False)
        priorites = collecte_composants_prioritaires(collecte)
        if matching is not None:
            reco = matching.rename(
                columns={
                    "Dataset ecoinvent": "Dataset ecoinvent reco",
                    "FE proposé (kg CO2e/kg)": "FE reco (kg CO2e/kg)",
                    "Géographie": "Géographie reco",
                    "Règle de matching": "Analyse / règle de matching",
                    "Statut": "Statut matching",
                }
            )
            reco = reco[
                [
                    "Component SKU",
                    "Dataset ecoinvent reco",
                    "FE reco (kg CO2e/kg)",
                    "Géographie reco",
                    "Analyse / règle de matching",
                    "Statut matching",
                ]
            ]
            priorites = priorites.merge(reco, on="Component SKU", how="left")
            priorites["A valider (OUI/NON)"] = None
        priorites.to_excel(writer, sheet_name="Priorites a investiguer", index=False)
        if matching is not None:
            matching.to_excel(writer, sheet_name="Matching ecoinvent", index=False)
        _style_collecte_sheet(writer.sheets["Fichier de collecte"], collecte)
        _style_synthese_sheet(writer.sheets["Synthese"], synthese)

